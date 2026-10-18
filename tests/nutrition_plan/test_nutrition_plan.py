"""Tests for scripts/nutrition_plan.py and SKILL.md.

The gate (the declared hard stops), the rules the targets apply, the workbook, and the check-in reader.
Every number here is invented. Run from the repository root:

    python3 -m unittest discover -s tests/nutrition_plan -t .

The build tests need openpyxl and are skipped without it.
"""
import datetime
import importlib.util
import json
import re
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

SKILL_DIR = Path(__file__).resolve().parents[2] / "plugins" / "ironledger" / "skills" / "ironledger-nutrition-plan"
SCRIPT = SKILL_DIR / "scripts" / "nutrition_plan.py"
SKILL = SKILL_DIR / "SKILL.md"
_spec = importlib.util.spec_from_file_location("nutrition_plan", SCRIPT)
np_ = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(np_)

try:
    import openpyxl  # noqa: F401
    HAVE_OPENPYXL = True
except ImportError:
    HAVE_OPENPYXL = False

# Every PMID the skill may cite. Each was checked on PubMed (first author, year, topic) on 2026-10-18.
VERIFIED_PMIDS = {
    "24092765", "24864135", "26817506", "28698222", "29497353", "23459753", "31247944", "37914977", "21558571",
    "34623696", "15883556", "7435418", "28682934", "41440795", "31895278", "26891166", "12519869", "37752011",
    "32915155", "32353079", "1615908", "1454084", "18198305", "21845193", "28925405", "33587549", "33630880",
    "29453741", "34070155", "28238618",
}


def spec(**over):
    s = {
        "client": {"name": "", "sex": "F", "age": 28, "height_cm": 168, "weight_kg": 65, "bodyfat_pct": None,
                   "activity": "moderate", "formula": "mifflin"},
        "flags": [],
        "phase": "Fat Loss Phase", "start_date": "2026-08-31", "weeks": 4,
        "targets": {"calories": 1750, "protein": 145, "carbs": 170, "fat": 55},
        "notes": ["Weigh meat, rice, pasta, oats and potato raw/dry on scales."],
        "days": [{"sheet": "Training Day", "meals": [{"name": "Meal 1", "items": [
            {"food": "oats", "qty": 50}, {"food": "chicken_breast", "qty": 200}]}]}],
    }
    for k, v in over.items():
        if k == "client":
            s["client"].update(v)
        else:
            s[k] = v
    return s


def targets(**kw):
    base = dict(sex="M", age=30, h=180, w=85, activity="moderate", goal="fat_loss")
    base.update(kw)
    return np_.targets(base.pop("sex"), base.pop("age"), base.pop("h"), base.pop("w"), base.pop("activity"),
                       base.pop("goal"), **base)


class Gate(unittest.TestCase):
    def refused(self, s, *words):
        with self.assertRaises(SystemExit) as cm:
            np_.gate(s)
        for w in words:
            self.assertIn(w, str(cm.exception))

    def test_a_clean_spec_passes_with_or_without_a_name(self):
        np_.gate(spec())
        np_.gate(spec(client={"name": "Test Person"}))

    def test_hard_stops_refuse_with_the_referral(self):  # r1
        for flag in ("pregnancy", "breastfeeding", "medical condition", "appetite medication", "weight medication",
                     "disordered eating", "RED-S"):
            with self.subTest(flag):
                self.refused(spec(flags=[flag]), "hard stop", "registered dietitian",
                             "Accredited Practising Dietitian")
        self.refused(spec(flags=["disordered eating"]), "1800 33 4673", "local eating disorder service")
        self.refused(spec(client={"age": 17}), "under 18")

    def test_the_age_is_required(self):
        self.refused(spec(client={"age": None}), "age")

    def test_other_flags_are_fine(self):
        np_.gate(spec(flags=["vegetarian", "lactose intolerance"]))

    def test_targets_refuse_hard_stops_too(self):
        with self.assertRaises(SystemExit):
            targets(age=16)
        with self.assertRaises(SystemExit):
            targets(flags=["pregnancy"])


class Targets(unittest.TestCase):
    def test_protein_on_a_cut_is_per_kg_of_fat_free_mass_when_measured(self):  # r9
        t = targets(bf=20)
        ffm = 85 * 0.8
        self.assertEqual(t["ffm"], ffm)
        self.assertEqual(t["protein"], round(2.5 * ffm / 5) * 5)
        self.assertTrue(2.3 * ffm <= t["protein"] <= 3.1 * ffm)
        lean = targets(bf=12)
        self.assertEqual(lean["protein"], round(2.8 * 85 * 0.88 / 5) * 5)
        self.assertTrue(any("g/kg FFM" in r for r in t["rules"]))

    def test_protein_without_body_fat_is_per_kg_of_bodyweight(self):
        self.assertEqual(targets()["protein"], round(2.2 * 85 / 5) * 5)

    def test_bmi_over_30_uses_the_reference_weight(self):
        t = targets(w=120)
        ref = round(27.5 * 1.8 ** 2, 1)
        self.assertEqual(t["ref_weight"], ref)
        self.assertEqual(t["protein"], round(2.2 * ref / 5) * 5)

    def test_calories_never_go_below_bmr(self):  # r8: no BMR x 0.8
        t = targets(activity="sedentary", rate=1.0, w=62, h=165, sex="F", age=45)
        self.assertGreaterEqual(t["calories"], t["bmr"] - 5)
        self.assertNotRegex(SCRIPT.read_text(), r"0\.8\s*\*\s*b\b|b\s*\*\s*0\.8")

    def test_energy_availability_is_flagged_under_30(self):  # r8
        t = targets(bf=18, exercise_kcal=900)
        self.assertIsNotNone(t["energy_availability"])
        self.assertTrue(any("under 30" in w for w in t["warnings"]))
        ok = targets(goal="maintain", bf=18, exercise_kcal=300)
        self.assertGreaterEqual(ok["energy_availability"], 30)
        self.assertFalse(any("under 30" in w for w in ok["warnings"]))
        self.assertTrue(any("--exercise-kcal" in w for w in targets(bf=18)["warnings"]))

    def test_rugby_in_season_sits_at_or_above_1_725(self):  # r4, r24
        t = targets(sport="rugby", season="in", goal="maintain")
        self.assertEqual(t["activity_factor"], 1.725)
        self.assertEqual(targets(sport="rugby", season="off", goal="maintain")["activity_factor"], 1.55)
        self.assertTrue(any("r24" in w for w in targets(sport="rugby", season="in")["warnings"]))

    def test_field_sport_and_adf_carbohydrate_floor(self):  # r12
        for sport in ("rugby", "field", "adf"):
            with self.subTest(sport):
                t = targets(sport=sport, activity="sedentary", rate=1.0, w=80, sex="F", age=35, h=165)
                self.assertGreaterEqual(t["carbs"], 3 * 80 - 5)
        plain = targets(activity="sedentary", rate=1.0, w=80, sex="F", age=35, h=165)
        self.assertLess(plain["carbs"], 3 * 80)

    def test_rest_days_keep_the_field_sport_floor(self):  # r12
        t = targets(sport="rugby", season="pre", sex="F", age=31, h=170, w=72.4, bf=27, rest_days=2)
        self.assertGreaterEqual(t["rest_day"]["carbs"], 3 * 72.4 - 5)
        self.assertGreaterEqual(t["training_day"]["carbs"], t["rest_day"]["carbs"])
        week = 5 * t["training_day"]["carbs"] + 2 * t["rest_day"]["carbs"]
        self.assertGreaterEqual(week, 7 * t["carbs"] - 10)

    def test_lean_and_performance_deficit_cap(self):  # r5, r6
        t = targets(w=100, bf=12, activity="hard")
        self.assertLessEqual(t["tdee"] - t["calories"], 505)
        self.assertIn("r5 fat-loss rate 0.6% BW/wk", targets(sport="powerlifting")["rules"])
        self.assertIn("r5 fat-loss rate 0.75% BW/wk", targets()["rules"])
        self.assertIn("r5 fat-loss rate 1.0% BW/wk", targets(w=120)["rules"])

    def test_fat_floor(self):  # r11
        t = targets(fat_pct=0.15)
        self.assertGreaterEqual(t["fat"] * 9 / t["calories"], 0.195)
        self.assertGreaterEqual(t["fat_gkg"], 0.5)
        self.assertTrue(any("20%" in w for w in t["warnings"]))
        peak = targets(fat_pct=0.15, peak=True)
        self.assertLess(peak["fat"] * 9 / peak["calories"], 0.2)

    def test_lean_gain_surplus_by_training_status(self):  # r7
        for status, pct in (("novice", 15), ("intermediate", 10), ("advanced", 5)):
            with self.subTest(status):
                t = targets(goal="lean_gain", status=status)
                self.assertAlmostEqual(t["calories"], round(t["tdee"] * (1 + pct / 100) / 10) * 10, delta=10)

    def test_fluid_and_fibre(self):  # r13
        self.assertEqual(targets(w=100)["fluid_l"], 3.3)
        self.assertEqual(targets(w=60, sex="F")["fluid_l"], 2.1)
        self.assertEqual(targets()["fibre_g"], 30)
        self.assertEqual(targets(sex="F")["fibre_g"], 25)

    def test_cunningham_needs_measured_body_fat(self):  # r3
        with self.assertRaises(SystemExit):
            targets(formula="cunningham")
        t = targets(formula="cunningham", bf=15)
        self.assertEqual(t["bmr"], round(500 + 22 * 85 * 0.85))


@unittest.skipUnless(HAVE_OPENPYXL, "build needs openpyxl")
class Workbook(unittest.TestCase):
    def test_the_name_is_optional_and_there_is_no_code(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            named = Path(tmp) / "Test Person - Fat Loss Phase Plan.xlsx"
            np_.build(spec(client={"name": "Test Person"}), str(named))
            wb = load_workbook(named)
            self.assertEqual(wb["Cover"]["B1"].value, "Test Person")
            self.assertEqual(wb["Cover"]["B2"].value, "Nutrition Plan")
            self.assertEqual(wb["Weekly check-in"]["B2"].value, "Test Person")
            unnamed = Path(tmp) / "plan.xlsx"
            np_.build(spec(), str(unnamed))
            wb = load_workbook(unnamed)
            self.assertEqual(wb["Cover"]["B1"].value, "Nutrition Plan")
            self.assertEqual(wb["Weekly check-in"]["B2"].value, "Weekly check-in")
            cells = [str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value]
            self.assertFalse(any(re.search(r"\bIC-\d", v) or "Coach" in v for v in cells))

    def test_the_sheets_and_the_check_in(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "plan.xlsx"
            np_.build(spec(weeks=3), str(out))
            wb = load_workbook(out)
        self.assertEqual(wb.sheetnames, ["Cover", "Training Day", "Weekly check-in", "Food list",
                                         "Free veg + 100cal fruit"])
        ws = wb["Weekly check-in"]
        banners = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value and str(ws.cell(r, 2).value).startswith("WEEK ")]
        self.assertEqual([ws.cell(r, 2).value for r in banners], ["WEEK 1", "WEEK 2", "WEEK 3"])
        days = [r for r in range(banners[0], banners[0] + 12) if ws.cell(r, 2).value in (
            "Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday")]
        self.assertEqual(len(days), 7)
        averages = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value == "Weekly average"]
        self.assertEqual(len(averages), 3)
        self.assertTrue(str(ws.cell(averages[0], 3).value).startswith("=IFERROR(ROUND(AVERAGE("))
        labels = {ws.cell(r, 6).value for r in range(1, ws.max_row + 1)}
        self.assertTrue({"Diet adherence", "Hunger", "Changes made"} <= labels)

    def test_build_refuses_a_hard_stop_and_writes_nothing(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "x.xlsx"
            with self.assertRaises(SystemExit):
                np_.build(spec(flags=["pregnancy"]), str(out))
            self.assertFalse(out.exists())


def week(start, mean, adherence=4, steps=9000, off_plan=0, **extra):
    w = {"start": start.isoformat(), "weights": [mean] * 7, "steps": [steps] * 7, "adherence": adherence,
         "off_plan_meals": off_plan, "sleep": 3, "energy": 3, "hunger": 3, "training": 3}
    w.update(extra)
    return w


def checkins(goal, rates, start_weight=80.0, plan_start=datetime.date(2026, 8, 31), weeks_before=0, last=None,
             **top):
    """A plan with one week per rate (% change from the week before), after `weeks_before` steady weeks."""
    weeks, mean, d = [], start_weight, plan_start
    for _ in range(weeks_before + 1):
        weeks.append(week(d, round(mean, 2)))
        d += datetime.timedelta(days=7)
    for i, r in enumerate(rates):
        mean = mean * (1 + r / 100)
        extra = last if (last and i == len(rates) - 1) else {}
        weeks.append(week(d, round(mean, 2), **extra))
        d += datetime.timedelta(days=7)
    data = {"goal": goal, "age": 30, "plan_start": plan_start.isoformat(), "as_of": (d - datetime.timedelta(days=1)).isoformat(),
            "weeks": weeks}
    data.update(top)
    return data


class Checkin(unittest.TestCase):
    def rules(self, data):
        return np_.checkin(data)["rules"]

    def test_a_slow_cut_lists_fewer_calories(self):  # r17
        out = np_.checkin(checkins("fat_loss", [-0.2, -0.3]))
        self.assertEqual(out["rules"], ["r17"])
        self.assertIn("fewer calories", out["options"][0]["change"])
        self.assertAlmostEqual(out["trend_pct_week"], -0.25, delta=0.02)
        self.assertEqual(len(out["weekly_rates"]), 2)

    def test_a_fast_cut_lists_more_calories(self):  # r18
        self.assertEqual(self.rules(checkins("fat_loss", [-1.2, -1.3])), ["r18"])
        self.assertEqual(self.rules(checkins("fat_loss", [-0.8, -0.9])), [])
        self.assertEqual(self.rules(checkins("fat_loss", [-0.8, -0.9], lean_or_performance=True)), ["r18"])

    def test_falling_performance_with_poor_sleep_lists_more_calories(self):  # r18
        out = self.rules(checkins("fat_loss", [-0.6, -0.6], last={"performance_falling": True, "sleep": 2}))
        self.assertEqual(out, ["r18"])

    def test_an_on_track_cut_changes_nothing(self):
        out = np_.checkin(checkins("fat_loss", [-0.7, -0.7]))
        self.assertEqual(out["rules"], [])
        self.assertEqual(out["options"][0]["change"], "no change: the trend is within the rules")

    def test_low_adherence_means_process_fixes_only(self):  # r15
        out = self.rules(checkins("fat_loss", [-0.2, -0.2], last={"adherence": 3}))
        self.assertEqual(out, ["r15"])
        self.assertEqual(self.rules(checkins("fat_loss", [-0.2, -0.2], last={"off_plan_meals": 2})), ["r15"])
        few = checkins("fat_loss", [-0.2, -0.2])
        few["weeks"][-1]["weights"] = [79.0, 79.0, 79.0, 79.0, None, None, None]
        self.assertEqual(self.rules(few), ["r15"])

    def test_steps_drop_comes_before_a_food_cut(self):  # r16
        data = checkins("fat_loss", [-0.2, -0.2], weeks_before=2, last={"steps": 7000})
        self.assertEqual(self.rules(data), ["r16", "r17"])
        self.assertLess(np_.checkin(data)["steps_vs_baseline_pct"], -15)

    def test_no_change_within_14_days_of_the_last_change(self):  # r14
        data = checkins("fat_loss", [-0.2, -0.2])
        data["last_change"] = (datetime.date.fromisoformat(data["as_of"]) - datetime.timedelta(days=10)).isoformat()
        out = np_.checkin(data)
        self.assertEqual(out["gates"]["wait"], "fail")
        self.assertEqual(out["rules"], [])
        self.assertIn("see the gates", out["options"][0]["change"])

    def test_a_decision_needs_three_weeks(self):  # r14
        out = np_.checkin(checkins("fat_loss", [-0.2]))
        self.assertEqual(out["weekly_rates"], [])
        self.assertIn("3 weeks", out["options"][0]["change"])
        one = np_.checkin(checkins("fat_loss", []))
        self.assertIn("the first week of the plan", one["gates"]["no_decision"])

    def test_a_no_decision_week_blocks_a_change(self):  # r21
        out = np_.checkin(checkins("fat_loss", [-0.2, -0.2], last={"no_decision": "travel"}))
        self.assertEqual(out["rules"], [])
        self.assertEqual(out["gates"]["no_decision"], "travel")

    def test_lean_gain_bands(self):  # r19
        self.assertEqual(np_.checkin(checkins("lean_gain", [0.1, 0.1]))["options"][0]["change"],
                         "+100-200 kcal, mostly carbohydrate")
        self.assertEqual(np_.checkin(checkins("lean_gain", [0.6, 0.6]))["options"][0]["change"], "-100-200 kcal")
        self.assertEqual(self.rules(checkins("lean_gain", [0.6, 0.6], status="novice")), [])
        self.assertEqual(self.rules(checkins("lean_gain", [0.3, 0.4])), [])

    def test_maintenance_drift(self):  # r20
        out = np_.checkin(checkins("maintain", [0.5, 0.4, 0.5]))
        self.assertEqual(out["rules"], ["r20"])
        self.assertTrue(out["options"][0]["change"].startswith("-100-200 kcal"))
        self.assertEqual(self.rules(checkins("maintain", [0.2, -0.2, 0.3])), [])
        self.assertEqual(self.rules(checkins("fat_loss", [-0.2, -0.2], in_season=True)), [])

    def test_a_long_cut_lists_a_diet_break(self):  # r22
        out = self.rules(checkins("fat_loss", [-0.7, -0.7], weeks_before=7))
        self.assertEqual(out, ["r22"])
        early = self.rules(checkins("fat_loss", [-0.7, -0.7], last={"hunger": 2, "adherence": 2, "training": 2}))
        self.assertIn("r22", early)

    def test_red_flags_mean_maintenance_and_a_doctor(self):  # r23
        out = np_.checkin(checkins("fat_loss", [-0.2, -0.2], red_flags=["missed period"]))
        self.assertEqual(out["rules"], ["r23"])
        self.assertIn("maintenance", out["options"][0]["change"])
        self.assertIn("doctor", out["options"][0]["change"])

    def test_hard_stops_refuse(self):  # r1
        for top in ({"flags": ["disordered eating"]}, {"age": 17}):
            with self.subTest(top):
                with self.assertRaises(SystemExit):
                    np_.checkin(checkins("fat_loss", [-0.2, -0.2], **top))

    def test_the_recheck_is_two_weeks_out(self):
        data = checkins("fat_loss", [-0.2, -0.2])
        out = np_.checkin(data)
        self.assertEqual(out["recheck"],
                         (datetime.date.fromisoformat(data["as_of"]) + datetime.timedelta(days=14)).isoformat())


class Cli(unittest.TestCase):
    def run_cli(self, *args):
        return subprocess.run([sys.executable, "-I", str(SCRIPT), *args], capture_output=True, text=True)

    def test_targets_runs_in_metric_and_imperial(self):
        r = self.run_cli("targets", "--sex", "F", "--age", "30", "--height", "165", "--weight", "70",
                         "--activity", "moderate", "--goal", "fat_loss", "--sport", "field", "--season", "pre")
        self.assertEqual(r.returncode, 0, r.stderr)
        metric = json.loads(r.stdout)
        self.assertIn("rules", metric)
        r = self.run_cli("targets", "--sex", "F", "--age", "30", "--height-in", "64.96", "--weight-lb", "154.32",
                         "--activity", "moderate", "--goal", "fat_loss", "--sport", "field", "--season", "pre")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(json.loads(r.stdout)["calories"], metric["calories"])

    def test_the_cut_command_is_gone(self):
        r = self.run_cli("cut", "--weight", "80", "--limit", "78", "--days-out", "7", "--weighin", "2h")
        self.assertNotEqual(r.returncode, 0)
        self.assertNotIn("def cut_plan", SCRIPT.read_text())

    def test_checkin_runs(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "checkins.json"
            p.write_text(json.dumps(checkins("fat_loss", [-0.2, -0.3])))
            r = self.run_cli("checkin", str(p))
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(json.loads(r.stdout)["rules"], ["r17"])

    def test_non_build_commands_need_only_the_standard_library(self):
        top = [l for l in SCRIPT.read_text().splitlines() if re.match(r"(import|from) ", l)]
        self.assertEqual(top, ["import json, sys, argparse, math, datetime, re"])


class SkillText(unittest.TestCase):
    def test_no_coaching_practice_left(self):
        text = SKILL.read_text() + SCRIPT.read_text()
        for word in ("screen_cleared", "screen: cleared", "FOR THE AI TOOL",
                     "Nutrition/Templates", "client code", "cut_plan"):
            with self.subTest(word):
                self.assertNotRegex(text, rf"\b{re.escape(word)}\b")
        self.assertNotRegex(text, r"\bIC-(<n>|\d)")
        self.assertNotIn(chr(0x2014), SKILL.read_text())

    def test_no_weight_cut_protocol(self):
        text = SKILL.read_text()
        for phrase in ("water load", "80-100 ml/kg", "5-15 ml/kg", "oral rehydration", "carb-up", "sweat suit"):
            with self.subTest(phrase):
                self.assertNotIn(phrase, text.lower())

    def test_every_pmid_was_verified(self):
        pmids = set(re.findall(r"PMID (\d+)", SKILL.read_text())) | set(re.findall(r"\band (\d{7,8})\b", SKILL.read_text()))
        self.assertGreater(len(pmids), 25)
        self.assertLessEqual(pmids, VERIFIED_PMIDS)

    def test_frontmatter(self):
        head = SKILL.read_text().split("---")[1]
        self.assertIn("name: ironledger-nutrition-plan", head)
        self.assertRegex(head, r"description: .{80,}")


if __name__ == "__main__":
    unittest.main()
