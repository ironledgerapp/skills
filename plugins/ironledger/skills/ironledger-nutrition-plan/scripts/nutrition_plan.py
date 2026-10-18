#!/usr/bin/env python3
"""IronLedger nutrition plan tool.

Sets calorie and macro targets from the evidence rules in SKILL.md (r1 to r25), builds an Excel meal plan with
like-for-like swaps and a weekly check-in sheet, and reads a lifter's own check-ins against the adjustment rules
(r14 to r23). Every command but `build` runs on the Python standard library; `build` needs openpyxl.

Commands
  targets  --sex M|F --age A (--height CM | --height-in IN) (--weight KG | --weight-lb LB) [--bf PCT]
           --activity LEVEL --goal GOAL [--rate PCT] [--surplus PCT] [--formula mifflin|cunningham|harris|katch]
           [--sport none|rugby|field|adf|powerlifting|other] [--season off|pre|in]
           [--status novice|intermediate|advanced] [--exercise-kcal K] [--flags "a,b"] [--peak] [--rest-days N]
  swap     FOOD QTY TO_FOOD [--basis protein|kcal|carbs]
  meal     food:qty food:qty ...           total an ad-hoc meal (use for OR/ alternative meals)
  foods    [CATEGORY]                      list food keys
  check    spec.json                       check day totals against targets
  build    spec.json out.xlsx              check, then build the workbook (needs openpyxl)
  checkin  checkins.json                   weekly averages, the trend and the r14 to r23 options
"""
import json, sys, argparse, math, datetime, re

__version__ = "1.0.0"

# ---------------------------------------------------------------------------
# FOOD DATABASE  (Australian supermarket foods)
# key: (display name, unit label, basis qty, kcal, protein, carbs, fat, category)
# basis qty = the amount the macros refer to (100 for g/ml foods, 1 for per-item foods)
# Sources: FSANZ AFCD Release 3 or Australian label panels (Woolworths/brand), checked Sep 2026.
# ---------------------------------------------------------------------------
FOODS = {
 # PROTEIN
 "chicken_breast":   ("Chicken breast (raw)", "g", 100, 99, 22.5, 0, 0.8, "PROTEIN"),
 "chicken_thigh":    ("Chicken thigh fillet, trimmed (raw)", "g", 100, 104, 19.1, 0, 3.0, "PROTEIN"),
 "beef_mince_5star": ("5 star beef mince (raw)", "g", 100, 129, 22.9, 0, 4.1, "PROTEIN"),
 "beef_mince_reg":   ("3-4 star beef mince (raw)", "g", 100, 170, 22.4, 0, 8.9, "PROTEIN"),
 "roo_mince":        ("Kangaroo mince (raw)", "g", 100, 103, 22.1, 0.3, 1.3, "PROTEIN"),
 "rump_steak":       ("Lean rump steak (raw)", "g", 100, 108, 20.4, 0, 2.8, "PROTEIN"),
 "pork_loin":        ("Lean pork loin steak (raw)", "g", 100, 112, 24.1, 0, 1.6, "PROTEIN"),
 "turkey_breast_mince": ("Turkey breast mince (raw)", "g", 100, 114, 23.7, 0, 1.5, "PROTEIN"),
 "turkey_mince":     ("Turkey mince, regular (raw)", "g", 100, 181, 17.6, 0, 12.3, "PROTEIN"),
 "prawns":           ("Prawns (raw)", "g", 100, 103, 23.4, 0, 0.9, "PROTEIN"),
 "salmon":           ("Atlantic salmon fillet (raw)", "g", 100, 231, 20.5, 0, 16.7, "PROTEIN"),
 "barramundi":       ("Barramundi fillet (raw)", "g", 100, 125, 19.2, 0, 5.3, "PROTEIN"),
 "basa":             ("Basa / white fish fillet (raw)", "g", 100, 79, 15.7, 0, 1.8, "PROTEIN"),
 "tuna_tin":         ("Tuna in springwater (95g tin)", "tin", 1, 70, 15.7, 0, 0.6, "PROTEIN"),
 "egg":              ("Egg (large)", "", 1, 74, 6.4, 0.7, 5.2, "PROTEIN"),
 "egg_white":        ("Liquid egg whites", "g", 100, 49, 11.6, 0.4, 0, "PROTEIN"),
 "ham":              ("Lean leg ham", "g", 100, 112, 17.0, 2.0, 2.5, "PROTEIN"),
 "tofu":             ("Firm tofu", "g", 100, 130, 12.8, 0, 8.3, "PROTEIN"),
 "jerky":            ("Beef jerky", "g", 100, 263, 34.0, 24.0, 3.0, "PROTEIN"),
 # DAIRY & ALTERNATIVES
 "milk_full":        ("Full cream milk", "ml", 100, 64, 3.3, 5.4, 3.4, "DAIRY"),
 "milk_lite":        ("Lite milk", "ml", 100, 43, 3.3, 4.8, 1.2, "DAIRY"),
 "milk_skim":        ("Skim milk", "ml", 100, 35, 3.6, 5.3, 0, "DAIRY"),
 "soy_lite":         ("Lite soy milk", "ml", 100, 39, 3.0, 2.5, 1.8, "DAIRY"),
 "oat_milk":         ("Oat milk (Vitasoy Oat Milky)", "ml", 100, 42, 0.5, 5.2, 2.0, "DAIRY"),
 "almond_milk":      ("Almond milk, unsweetened", "ml", 100, 17, 0.6, 0.3, 1.4, "DAIRY"),
 "greek_yog_plain":  ("Chobani plain Greek yoghurt (0% fat)", "g", 100, 57, 9.7, 4.1, 0.2, "DAIRY"),
 "greek_yog_flav":   ("Chobani flavoured Greek yoghurt", "g", 100, 75, 8.1, 9.9, 0.2, "DAIRY"),
 "yopro":            ("YoPRO high protein yoghurt", "g", 100, 60, 10.6, 3.5, 0.3, "DAIRY"),
 "cheese_tasty":     ("Tasty cheese", "g", 100, 411, 24.8, 1.0, 34.7, "DAIRY"),
 "cheese_light":     ("Light tasty cheese", "g", 100, 324, 27.9, 0, 23.4, "DAIRY"),
 "cottage_cheese":   ("Light cottage cheese", "g", 100, 93, 11.3, 6.3, 2.4, "DAIRY"),
 # CARBS
 "rice_dry":         ("White rice (dry weight)", "g", 100, 340, 7.0, 74.0, 0.9, "CARBS"),
 "pasta_dry":        ("Pasta (dry weight)", "g", 100, 349, 12.7, 68.5, 1.4, "CARBS"),
 "rice_cooked":      ("White rice (cooked)", "g", 100, 158, 3.1, 34.7, 0.2, "CARBS"),
 "pasta_cooked":     ("Pasta (cooked)", "g", 100, 137, 5.3, 26.2, 0.4, "CARBS"),
 "microwave_rice":   ("Microwave rice pouch (250g)", "pouch", 1, 418, 9.6, 85.2, 3.0, "CARBS"),
 "potato":           ("Potato (raw)", "g", 100, 58, 2.2, 11.0, 0, "CARBS"),
 "sweet_potato":     ("Sweet potato (raw)", "g", 100, 65, 1.4, 13.2, 0.2, "CARBS"),
 "oats":             ("Rolled oats", "g", 100, 382, 12.8, 56.7, 9.2, "CARBS"),
 "weetbix":          ("Weet-Bix", "", 1, 55, 1.9, 10.2, 0.2, "CARBS"),
 "muesli":           ("Toasted muesli", "g", 100, 447, 10.7, 61.0, 15.7, "CARBS"),
 "just_right":       ("Kellogg's Just Right", "g", 100, 363, 8.1, 70.9, 2.6, "CARBS"),
 "sourdough":        ("Sourdough bread", "g", 100, 229, 8.6, 44.0, 1.4, "CARBS"),
 "bread_white":      ("White sandwich bread (slice)", "", 1, 86, 3.4, 15.8, 1.0, "CARBS"),
 "bread_wholemeal":  ("Wholemeal bread (slice)", "", 1, 97, 3.6, 16.0, 1.3, "CARBS"),
 "english_muffin":   ("English muffin", "", 1, 132, 4.8, 23.5, 1.4, "CARBS"),
 "wrap":             ("Mission Original wrap", "", 1, 209, 5.8, 32.3, 5.8, "CARBS"),
 "sandwich_thin":    ("Tip Top Sandwich Thin", "", 1, 100, 3.8, 17.2, 1.0, "CARBS"),
 "corn_thin":        ("Corn Thins (slice)", "", 1, 23, 0.6, 4.3, 0.2, "CARBS"),
 "rice_cake_thick":  ("Thick rice cake", "", 1, 49, 1.0, 9.8, 0.4, "CARBS"),
 "rice_crackers":    ("Rice crackers (Sakata plain)", "g", 100, 409, 7.5, 87.1, 2.9, "CARBS"),
 "chickpeas":        ("Chickpeas, canned (drained)", "g", 100, 113, 6.3, 13.3, 2.1, "CARBS"),
 "four_bean":        ("Four bean mix, canned", "g", 100, 126, 8.1, 16.5, 1.6, "CARBS"),
 "lentils_dry":      ("Red lentils (dry)", "g", 100, 326, 23.0, 45.7, 1.9, "CARBS"),
 # FRUIT
 "fruit100":         ("Banana OR 100cal fruit", "", 1, 100, 1.5, 21.0, 0.2, "FRUIT"),
 "banana":           ("Banana (medium, 120g flesh)", "", 1, 114, 1.7, 23.9, 0.2, "FRUIT"),
 "apple":            ("Apple", "g", 100, 58, 0.2, 12.7, 0, "FRUIT"),
 "orange":           ("Orange", "g", 100, 42, 1.0, 8.2, 0, "FRUIT"),
 "berries_frozen":   ("Frozen mixed berries", "g", 100, 48, 0.6, 9.6, 0.1, "FRUIT"),
 "blueberries":      ("Blueberries", "g", 100, 49, 0.5, 10.4, 0, "FRUIT"),
 "pineapple":        ("Pineapple", "g", 100, 39, 0.5, 8.6, 0, "FRUIT"),
 "orange_juice":     ("Orange juice", "ml", 100, 42, 1.0, 10.2, 0, "FRUIT"),
 # VEG
 "free_veg":         ("Free veggies", "g", 100, 0, 0, 0, 0, "VEG"),
 "pumpkin":          ("Pumpkin (raw)", "g", 100, 34, 1.2, 6.7, 0, "VEG"),
 "frozen_mixed_veg": ("Frozen peas, corn & carrot", "g", 100, 63, 3.4, 9.5, 0.5, "VEG"),
 "avocado":          ("Avocado", "g", 100, 156, 1.6, 2.2, 14.6, "VEG"),
 # FATS & CONDIMENTS
 "butter":           ("Butter", "g", 100, 734, 1.1, 0.6, 82.2, "FATS & SAUCES"),
 "olive_oil":        ("Olive oil", "ml", 100, 883, 0, 0, 99.9, "FATS & SAUCES"),
 "pb":               ("Natural peanut butter", "g", 100, 581, 25.0, 13.4, 49.3, "FATS & SAUCES"),
 "almonds":          ("Almonds", "g", 100, 570, 19.7, 5.4, 50.5, "FATS & SAUCES"),
 "sauce":            ("Sauce of choice (tomato/BBQ/sweet chilli)", "ml", 100, 157, 0.5, 37.4, 0.3, "FATS & SAUCES"),
 "mayo_ff":          ("99% fat free mayo", "g", 100, 115, 1.0, 24.9, 1.0, "FATS & SAUCES"),
 "perinaise":        ("Nando's Perinaise", "g", 100, 223, 0.5, 12.5, 19.1, "FATS & SAUCES"),
 "pasta_sauce":      ("Tomato pasta sauce", "g", 100, 38, 1.1, 7.3, 0.2, "FATS & SAUCES"),
 "honey":            ("Honey", "g", 100, 339, 0.3, 83.1, 0, "FATS & SAUCES"),
 "jam":              ("Jam", "g", 100, 263, 0.3, 67.3, 0, "FATS & SAUCES"),
 # SUPPLEMENTS
 "wpi":              ("Whey protein isolate (30g scoop)", "scoop", 1, 113, 26.4, 1.0, 0.3, "SUPPLEMENTS"),
 "whey":             ("Protein powder (30g scoop)", "scoop", 1, 120, 22.2, 3.1, 1.9, "SUPPLEMENTS"),
 # SNACKS & SPORTS FOODS
 "lcm":              ("LCMs bar", "", 1, 85, 0.7, 15.3, 2.3, "SNACKS"),
 "milo_bar":         ("Milo snack bar", "", 1, 72, 1.1, 12.5, 1.0, "SNACKS"),
 "snake":            ("Allen's snake (lolly)", "", 1, 42, 0.6, 9.6, 0, "SNACKS"),
 "timtam":           ("Tim Tam", "", 1, 96, 0.8, 11.9, 4.9, "SNACKS"),
 "cadbury":          ("Cadbury Dairy Milk", "g", 100, 535, 7.4, 57.2, 30.5, "SNACKS"),
 "chips_28g":        ("Red Rock Deli chips (28g bag)", "", 1, 136, 2.2, 16.9, 6.4, "SNACKS"),
 "upgo":             ("Up&Go original (250ml)", "", 1, 196, 8.3, 28.5, 4.3, "SNACKS"),
 "upgo_protein":     ("Up&Go Protein Energize (250ml)", "", 1, 217, 17.5, 25.0, 4.5, "SNACKS"),
 "sports_drink":     ("Sports drink (Gatorade/Maximus)", "ml", 100, 28, 0, 6.8, 0, "SNACKS"),
}

ACTIVITY = {"sedentary": 1.2, "light": 1.375, "moderate": 1.55, "hard": 1.725, "extreme": 1.9}

def food(key):
    if key not in FOODS:
        raise SystemExit(f"Unknown food key '{key}'. Run: foods")
    return FOODS[key]

def macros(key, qty):
    n, u, basis, k, p, c, f, cat = food(key)
    s = qty / basis
    return k * s, p * s, c * s, f * s

# ---------------------------------------------------------------------------
# GATE (r1): the hard stops the lifter declares. Any one means no calorie target and a referral instead.
# ---------------------------------------------------------------------------
HARD_STOPS = ("pregnancy", "breastfeeding", "medical condition", "appetite medication", "weight medication",
              "disordered eating", "red-s")
REFERRAL = ("No plan. See a doctor or a registered dietitian (in Australia, an Accredited Practising Dietitian: "
            "dietitiansaustralia.org.au). For disordered eating, also contact the Butterfly helpline on "
            "1800 33 4673 (Australia) or your local eating disorder service.")


def hard_stops(age=None, flags=()):
    """The reasons a plan can't be written, from the age and the declared hard-stop flags. Empty means none."""
    out = [f for f in (str(x).strip().lower() for x in flags) if f in HARD_STOPS]
    if age is not None and float(age) < 18:
        out.append("under 18")
    return out


def gate(spec):
    """Refuse a spec with a declared hard stop or an age under 18, or one that gives no age."""
    c = spec.get("client", {})
    if c.get("age") is None:
        raise SystemExit("plan refused: client.age is missing; the plan needs the lifter's age (r1)")
    stops = hard_stops(c.get("age"), spec.get("flags", []))
    if stops:
        raise SystemExit("plan refused: hard stop (" + ", ".join(stops) + "). " + REFERRAL)


# ---------------------------------------------------------------------------
# TARGETS (rules r3 to r13, r24 and r25; each output lists the rules it applied)
# ---------------------------------------------------------------------------
FIELD_SPORTS = ("rugby", "field")           # team or field sport: carbohydrate never under 3 g/kg (r12)
PERFORMANCE_SPORTS = FIELD_SPORTS + ("adf", "powerlifting")
NRV_FLUID_L = {"M": 2.6, "F": 2.1}          # NRV adequate intake, total water from drinks (r13)
FIBRE_G = {"M": 30, "F": 25}                # NRV adequate intake (r13)


def lean(male, bf):
    return bf is not None and ((male and bf < 15) or (not male and bf < 23))


def bmr(sex, age, h, w, bf=None, formula="mifflin"):
    male = sex.upper().startswith("M")
    if formula in ("katch", "cunningham"):
        if bf is None:
            raise SystemExit(f"{formula} needs a measured body fat % (--bf)")
        ffm = w * (1 - bf / 100)
        # Cunningham 1980 (r3): only with DXA or trained-skinfold FFM in trained lifters. Katch-McArdle is kept
        # for old plans only; it isn't in the evidence list.
        return 500 + 22 * ffm if formula == "cunningham" else 370 + 21.6 * ffm
    if formula == "harris":
        return (66.473 + 13.7516 * w + 5.0033 * h - 6.755 * age) if male else (655.0955 + 9.5634 * w + 1.8496 * h - 4.6756 * age)
    return 10 * w + 6.25 * h - 5 * age + (5 if male else -161)


def targets(sex, age, h, w, activity, goal, bf=None, formula="mifflin", rate=None, surplus=None, protein_gkg=None,
            fat_pct=None, rest_days=0, sport="none", season=None, status="intermediate", exercise_kcal=None,
            flags=(), peak=False):
    stops = hard_stops(age, flags)
    if stops:
        raise SystemExit("targets refused: hard stop (" + ", ".join(stops) + "). " + REFERRAL)
    male = sex.upper().startswith("M")
    rules, warn = [], []
    b = bmr(sex, age, h, w, bf, formula)
    rules.append(f"r3 BMR {formula}")
    factor = ACTIVITY[activity]
    if sport in FIELD_SPORTS and season == "in" and factor < ACTIVITY["hard"]:
        factor = ACTIVITY["hard"]                                  # r4: rugby in season at least 1.725
        rules.append("r4 in-season field sport: activity factor raised to 1.725 (consider 1.9)")
    tdee = b * factor
    bmi = w / (h / 100) ** 2
    ref_w = w if bmi <= 30 else round(27.5 * (h / 100) ** 2, 1)   # protein/fat reference weight for high BMI (r9)
    ffm = round(w * (1 - bf / 100), 1) if bf is not None else None
    is_lean = lean(male, bf)
    performance = sport in PERFORMANCE_SPORTS or goal == "performance"
    protein_basis = f"{'reference weight' if ref_w != w else 'bodyweight'}"
    if goal == "fat_loss":
        if rate is None:                                           # r5
            if bmi > 30 or (bf is not None and ((male and bf > 25) or (not male and bf > 32))):
                rate = 1.0
            elif is_lean:
                rate = 0.5
            elif performance or status == "advanced":
                rate = 0.6
            else:
                rate = 0.75
        rules.append(f"r5 fat-loss rate {rate}% BW/wk")
        if sport in FIELD_SPORTS and season == "in":               # r24: in season sits at maintenance
            warn.append("In-season field sport normally sits at maintenance (r24): weigh a fat-loss goal now against performance")
        deficit = w * rate / 100 * 7700 / 7
        cap = 0.25 * tdee
        if (is_lean or performance) and status != "novice":
            cap = min(cap, 500)                                    # r6: lean trained lifters ~500 kcal/day at most
        if deficit > cap:
            deficit = cap
            rules.append(f"r6 deficit capped at {round(cap)} kcal/day")
        kcal = tdee - deficit
        if ffm is not None and protein_gkg is None:                # r9: 2.3-3.1 g/kg FFM in a deficit
            pg_ffm = 2.8 if is_lean else 2.5
            protein = round(pg_ffm * ffm / 5) * 5
            protein_basis = f"{pg_ffm} g/kg FFM ({ffm} kg FFM)"
        else:
            protein = round((protein_gkg or 2.2) * ref_w / 5) * 5  # FFM unknown: about 2.0-2.4 g/kg BW
        fp = fat_pct or 0.25
    elif goal == "lean_gain":                                      # r7
        default = {"novice": 15, "intermediate": 10, "advanced": 5}.get(status, 10)
        sp = (surplus if surplus is not None else default) / 100
        rules.append(f"r7 surplus {round(sp * 100)}% ({status})")
        kcal = tdee * (1 + sp)
        protein = round((protein_gkg or 1.8) * ref_w / 5) * 5
        fp = fat_pct or 0.25
    elif goal in ("maintain", "recomp", "performance"):
        kcal = tdee
        protein = round((protein_gkg or (2.0 if goal != "maintain" else 1.8)) * ref_w / 5) * 5
        fp = fat_pct or 0.25
    else:
        raise SystemExit("goal must be fat_loss | lean_gain | maintain | recomp | performance")
    rules.append(f"r9 protein {protein_basis}")
    if kcal < b:                                                   # r8: never below estimated BMR (no BMR x 0.8)
        kcal = b
        rules.append("r8 calories raised to estimated BMR (the floor); the expected rate is slower")
        warn.append("Calories were below estimated BMR and have been raised to it: expect a slower rate")
    if fp < 0.20 and not peak:                                     # r11: fat at least 20% of energy
        fp = 0.20
        warn.append("Fat raised to 20% of energy (15-20% only briefly, for men, in a peak phase: --peak)")
    kcal = round(kcal / 10) * 10
    fat = max(fp * kcal / 9, 0.5 * ref_w)                          # r11: and at least 0.5 g/kg
    fat = round(fat / 5) * 5
    rules.append(f"r11 fat {round(fat * 9 / kcal * 100)}% of energy")
    carbs = round((kcal - protein * 4 - fat * 9) / 4 / 5) * 5
    if sport in FIELD_SPORTS + ("adf",) and carbs < 3 * ref_w:     # r12: field-sport and ADF lifters never under 3 g/kg
        carbs = round(3 * ref_w / 5) * 5
        kcal = protein * 4 + carbs * 4 + fat * 9
        rules.append("r12 carbohydrate raised to the 3 g/kg floor for a field-sport or ADF candidate; calories follow")
    ckg = carbs / ref_w
    if ckg < 2.5 and activity in ("moderate", "hard", "extreme") and sport not in FIELD_SPORTS + ("adf",):
        warn.append(f"Carbs only {ckg:.1f} g/kg: low for this training load, consider a slower rate")
    exp_kg = (kcal - tdee) * 7 / 7700
    out = dict(bmr=round(b), tdee=round(tdee), activity_factor=factor, bmi=round(bmi, 1), ref_weight=ref_w, ffm=ffm,
               calories=kcal, protein=protein, carbs=carbs, fat=fat, protein_gkg=round(protein / w, 2),
               carbs_gkg=round(carbs / w, 2), fat_gkg=round(fat / w, 2), expected_kg_per_wk=round(exp_kg, 2),
               expected_pct_per_wk=round(exp_kg / w * 100, 2),
               fluid_l=round(max(NRV_FLUID_L["M" if male else "F"], 0.033 * w), 1),
               fibre_g=FIBRE_G["M" if male else "F"], rules=rules, warnings=warn)
    out["fluid_note"] = (f"Water at least {out['fluid_l']} L a day, plus 500-750 ml per training hour (r13)")
    if ffm is not None and exercise_kcal is not None:              # r8: energy availability
        ea = (kcal - exercise_kcal) / ffm
        out["energy_availability"] = round(ea, 1)
        if ea < 30:
            warn.append(f"Energy availability {ea:.0f} kcal/kg FFM/day is under 30 (r8): don't start this target on "
                        "your own; eat more or train less, or see a sports dietitian or doctor")
    elif ffm is not None:
        out["energy_availability"] = None
        warn.append("Energy availability not computed: give --exercise-kcal (average daily training energy)")
    if goal == "lean_gain":
        lo, hi = (0.1, 0.25) if status == "advanced" else (0.25, 0.75 if status == "novice" else 0.5)
        if not (lo <= out["expected_pct_per_wk"] <= hi):
            warn.append(f"Expected gain outside {lo}-{hi}%/wk: adjust --surplus")
    if rest_days and rest_days >= 2:
        td = 7 - rest_days
        ct = round(7 * carbs / (td + 0.8 * rest_days) / 5) * 5
        cr = round(0.8 * ct / 5) * 5
        floor = round(3 * ref_w / 5) * 5
        if sport in FIELD_SPORTS + ("adf",) and cr < floor:        # r12 holds on rest days too
            cr = floor
            ct = max(floor, round((7 * carbs - rest_days * cr) / td / 5) * 5)
            rules.append("r12 rest-day carbohydrate held at 3 g/kg; training days carry the rest")
        out["training_day"] = dict(calories=protein * 4 + ct * 4 + fat * 9, protein=protein, carbs=ct, fat=fat)
        out["rest_day"] = dict(calories=protein * 4 + cr * 4 + fat * 9, protein=protein, carbs=cr, fat=fat)
        out["split_note"] = f"{td} training days + {rest_days} rest days average {round((td*out['training_day']['calories']+rest_days*out['rest_day']['calories'])/7)} kcal"
    elif rest_days:
        out["split_note"] = "Fewer than 2 rest days: use one daily plan"
    return out


# ---------------------------------------------------------------------------
# SWAPS
# ---------------------------------------------------------------------------
ROUND = {"g": 5, "ml": 10}
def swap_qty(from_key, qty, to_key, basis="auto"):
    fk, fp, fc, ff = macros(from_key, qty)
    tn, tu, tb, tk, tp, tc, tf, tcat = food(to_key)
    if basis == "auto":
        basis = "protein" if food(from_key)[7] == "PROTEIN" else "kcal"
    src = {"protein": fp, "carbs": fc, "kcal": fk}[basis]
    per = {"protein": tp, "carbs": tc, "kcal": tk}[basis] / tb
    q = src / per
    step = ROUND.get(tu, 0.5)
    q = round(q / step) * step
    if tu in ("g", "ml") and q >= 100:
        q = round(q / 10) * 10
    return q

def fmt_qty(key, q):
    n, u, *_ = food(key)
    qs = f"{q:g}"
    if u in ("g", "ml"):
        return f"{qs}{u} {n}"
    if u:
        return f"{qs} {u}{('es' if u.endswith('ch') else 's') if q != 1 and not u.endswith('s') else ''} {n}"
    return f"{qs}x {n}"

# ---------------------------------------------------------------------------
# WORKBOOK
# ---------------------------------------------------------------------------
def build(spec, out):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    BLUE = PatternFill("solid", fgColor="66CCFF")
    thin, thick = Side(style="thin"), Side(style="thick")
    B_THIN = Border(left=thin, right=thin, top=thin, bottom=thin)
    B_HEAD = Border(left=thin, right=thin, top=thick, bottom=thick)
    F = lambda sz=10, b=False: Font(name="Calibri", size=sz, bold=b)
    C = Alignment(horizontal="center", vertical="center")
    L = Alignment(horizontal="left", vertical="center", wrap_text=False)

    gate(spec)
    c = spec["client"]
    name = (c.get("name") or "").strip()                         # optional: the lifter's name on the plan
    title = spec.get("title", "Nutrition Plan")
    wb = Workbook()

    # ---------------- Cover ----------------
    ws = wb.active; ws.title = "Cover"
    for col, wdt in zip("ABCDEF", (3, 18, 20, 4, 18, 14)):
        ws.column_dimensions[col].width = wdt
    ws["B1"] = name or title; ws["B1"].font = F(36, True)
    if name:
        ws["B2"] = title; ws["B2"].font = F(20, True)
    ws["B3"] = spec.get("phase", ""); ws["B3"].font = F(14, True)
    rows = [("Gender", "Male" if c["sex"].upper().startswith("M") else "Female"),
            ("Formula", {"mifflin": "Mifflin St Jeor", "cunningham": "Cunningham", "harris": "Harris-Benedict", "katch": "Katch-McArdle"}[c.get("formula", "mifflin")]),
            ("Age", c["age"]), ("Height (cm)", c["height_cm"]), ("Weight (kg)", c["weight_kg"]),
            ("Activity", c["activity"]), ("Body fat %", c.get("bodyfat_pct")), ("LBM (kg)", None)]
    r0 = 6
    for i, (k, v) in enumerate(rows):
        r = r0 + i
        ws.cell(r, 2, k).font = F(14, True)
        cell = ws.cell(r, 3, v); cell.font = F(14)
    # rows: 6 Gender, 7 Formula, 8 Age, 9 Height, 10 Weight, 11 Activity, 12 Body fat %, 13 LBM
    ws["C13"] = '=IF(C12="","",ROUND(C10*(1-C12/100),1))'
    ws["E6"] = "BMR"; ws["E7"] = "TDEE"; ws["E9"] = "Daily target"; 
    ws["F6"] = ('=ROUND(IF(C7="Cunningham",500+22*C13,IF(C7="Katch-McArdle",370+21.6*C13,IF(C7="Harris-Benedict",'
                'IF(C6="Male",66.473+13.7516*C10+5.0033*C9-6.755*C8,655.0955+9.5634*C10+1.8496*C9-4.6756*C8),'
                '10*C10+6.25*C9-5*C8+IF(C6="Male",5,-161)))),0)')
    ws["F7"] = ('=ROUND(F6*IF(C11="sedentary",1.2,IF(C11="light",1.375,IF(C11="moderate",1.55,IF(C11="hard",1.725,1.9)))),0)')
    t = spec.get("targets", {})
    trows = [("Calories", t.get("calories")), ("Protein (g)", t.get("protein")), ("Carbs (g)", t.get("carbs")), ("Fat (g)", t.get("fat"))]
    for i, (k, v) in enumerate(trows):
        ws.cell(10 + i, 5, k).font = F(14, True); ws.cell(10 + i, 6, v).font = F(14)
    for r in (6, 7):
        ws.cell(r, 5).font = F(14, True); ws.cell(r, 6).font = F(14)
    ws["E9"].font = F(14, True)
    if spec.get("start_date"):
        ws["E15"] = "Plan start"; ws["F15"] = datetime.date.fromisoformat(spec["start_date"]); ws["F15"].number_format = "dd/mm/yyyy"
        ws["E15"].font = F(14, True); ws["F15"].font = F(14)
    if spec.get("event_date"):
        ws["E16"] = spec.get("event_name") or "Event date"; ws["F16"] = datetime.date.fromisoformat(spec["event_date"]); ws["F16"].number_format = "dd/mm/yyyy"
        ws["E16"].font = F(14, True); ws["F16"].font = F(14)
    notes = spec.get("notes") or []
    if notes:
        ws["B16"] = "Plan rules"; ws["B16"].font = F(14, True)
        for i, n in enumerate(notes):
            ws.cell(17 + i, 2, f"{i+1}. {n}").font = F(11)
    dr = 18 + len(notes) + 1
    ws.cell(dr, 2, "Day plans").font = F(14, True)
    for j, h in enumerate(("Calories", "Protein", "Carbs", "Fat")):
        ws.cell(dr, 3 + j + (1 if j >= 1 else 0), h).font = F(11, True)
    for i, day in enumerate(spec["days"]):
        sn = day["sheet"][:31]
        ws.cell(dr + 1 + i, 2, sn).font = F(11)
        for j, col in enumerate("EFGH"):
            cell = ws.cell(dr + 1 + i, 3 + j + (1 if j >= 1 else 0), f"='{sn}'!{col}4"); cell.number_format = "0"; cell.font = F(11)

    # ---------------- Meal plan sheets ----------------
    first_day_total = None
    for day in spec["days"]:
        ws = wb.create_sheet(day["sheet"][:31])
        for col, wdt in zip("ABCDEFGH", (2.7, 6.5, 4, 46, 9, 9, 9, 9)):
            ws.column_dimensions[col].width = wdt
        ws.merge_cells("B2:F2"); ws.merge_cells("G2:H2"); ws.merge_cells("B3:D4")
        ws["B2"] = day.get("title", spec.get("phase", "")); ws["B3"] = day.get("subtitle", day["sheet"])
        for rng in ("B2", "G2", "B3"):
            ws[rng].font = F(18, True)
        for rr in range(2, 5):
            for cc in range(2, 9):
                ws.cell(rr, cc).fill = BLUE
        ws["B2"].alignment = L; ws["G2"].alignment = Alignment(horizontal="right"); ws["B3"].alignment = L
        for j, h in enumerate(("Calories", "Protein", "Carbs", "Fat")):
            cell = ws.cell(3, 5 + j, h); cell.font = F(10, True); cell.alignment = C; cell.border = B_THIN
        r = 6
        meal_total_rows = []
        for meal in day["meals"]:
            ws.merge_cells(start_row=r, start_column=2, end_row=r + 1, end_column=4)
            ws.cell(r, 2, meal["name"]).font = F(13, True)
            ws.cell(r, 2).alignment = L
            for rr in (r, r + 1):
                for cc in range(2, 9):
                    ws.cell(rr, cc).fill = BLUE
            for j, h in enumerate(("Calories", "Protein", "Carbs", "Fat")):
                cell = ws.cell(r, 5 + j, h); cell.font = F(10, True); cell.alignment = C; cell.border = B_THIN
            tot_row = r + 1
            meal_total_rows.append(tot_row)
            first_item = r + 2
            rr = r + 3  # leave one blank row like the template
            for it in meal["items"]:
                key, qty = it["food"], it["qty"]
                n, u, basis, k, p, cc_, f, cat = food(key)
                label = it.get("label") or n
                ws.cell(rr, 2, qty).font = F(); ws.cell(rr, 2).alignment = Alignment(horizontal="right"); ws.cell(rr, 2).number_format = "General"
                ws.cell(rr, 3, u).font = F()
                ws.cell(rr, 4, label).font = F()
                for j, v in enumerate((k, p, cc_, f)):
                    coef = v / basis
                    cell = ws.cell(rr, 5 + j, f"={coef:g}*B{rr}" if v else 0)
                    cell.font = F(); cell.alignment = C; cell.number_format = "0"; cell.border = B_THIN
                for cidx in (2, 3, 4):
                    ws.cell(rr, cidx).border = B_THIN
                rr += 1
                for s in it.get("swaps", []):
                    cell = ws.cell(rr, 4, s if s.startswith("OR/") else "OR/ " + s); cell.font = F(); cell.alignment = L
                    rr += 1
            for a in meal.get("alts", []):
                rr += 1
                cell = ws.cell(rr, 2, a if a.startswith("OR/") else "OR/ " + a); cell.font = F(); cell.alignment = L
                rr += 1
            last_item = rr
            for j in range(4):
                col = get_column_letter(5 + j)
                cell = ws.cell(tot_row, 5 + j, f"=SUM({col}{first_item}:{col}{last_item})")
                cell.font = F(10, True); cell.alignment = C; cell.number_format = "0"; cell.border = B_HEAD
            r = rr + 2
        for j in range(4):
            col = get_column_letter(5 + j)
            cell = ws.cell(4, 5 + j, "=SUM(" + ",".join(f"{col}{x}" for x in meal_total_rows) + ")")
            cell.font = F(10, True); cell.alignment = C; cell.number_format = "0"; cell.border = B_HEAD
        ws["G2"] = "=E4"; ws["G2"].number_format = "0"
        if first_day_total is None:
            first_day_total = day["sheet"][:31]

    # ---------------- Weekly check-in ----------------
    weeks = int(spec.get("weeks", 8))
    ws = wb.create_sheet("Weekly check-in")
    for col, wdt in zip("ABCDEFGHIJ", (2.7, 18, 11, 11, 3, 24, 8, 60, 3, 3)):
        ws.column_dimensions[col].width = wdt
    ws["B2"] = name or "Weekly check-in"; ws["B2"].font = F(18, True)
    ws["F2"] = "Nutrition phase:"; ws["G2"] = spec.get("phase", ""); ws["F2"].font = F(11, True)
    ws["F3"] = "Plan start:"; ws["F3"].font = F(11, True)
    if spec.get("start_date"):
        ws["G3"] = datetime.date.fromisoformat(spec["start_date"]); ws["G3"].number_format = "dd/mm/yyyy"
    ws["B4"] = "Weigh daily: morning, after toilet, before food. Ratings 1 (poor) to 5 (great)."; ws["B4"].font = F(10)
    q = f"'{first_day_total}'!"
    r = 6
    prev_avg = None
    ratings = ["Training performance", "Training recovery", "Diet adherence", "Sleep quality", "General energy levels", "Hunger"]
    for wk in range(1, weeks + 1):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        ws.cell(r, 2, f"WEEK {wk}").font = F(14, True)
        for cc in range(2, 9):
            ws.cell(r, cc).fill = BLUE
        hdr = r + 1
        ws.cell(hdr, 2, "=G3+7*%d" % (wk - 1) if spec.get("start_date") else "Week start").number_format = "ddd dd/mm"
        ws.cell(hdr, 3, "Weight"); ws.cell(hdr, 4, "Steps"); ws.cell(hdr, 6, "Check-in"); ws.cell(hdr, 7, "1-5"); ws.cell(hdr, 8, "Comments")
        for cc in (2, 3, 4, 6, 7, 8):
            ws.cell(hdr, cc).font = F(10, True); ws.cell(hdr, cc).border = B_THIN
        wd = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
        s0 = datetime.date.fromisoformat(spec["start_date"]).weekday() if spec.get("start_date") else 0
        for i, dname in enumerate(wd[s0:] + wd[:s0]):
            rr = hdr + 1 + i
            ws.cell(rr, 2, dname).font = F()
            for cc in (2, 3, 4):
                ws.cell(rr, cc).border = B_THIN
        for i, lab in enumerate(ratings):
            rr = hdr + 1 + i
            ws.cell(rr, 6, lab).font = F(); 
            for cc in (6, 7, 8):
                ws.cell(rr, cc).border = B_THIN
        avg = hdr + 8
        ws.cell(avg, 2, "Weekly average").font = F(10, True)
        ws.cell(avg, 3, f"=IFERROR(ROUND(AVERAGE(C{hdr+1}:C{hdr+7}),1),\"\")").font = F(10, True)
        ws.cell(avg, 4, f"=IFERROR(ROUND(AVERAGE(D{hdr+1}:D{hdr+7}),0),\"\")").font = F(10, True)
        if prev_avg:
            ws.cell(avg + 1, 2, "Weight change (kg)").font = F(10, True)
            ws.cell(avg + 1, 3, f"=IFERROR(C{avg}-C{prev_avg},\"\")")
            ws.cell(avg + 2, 2, "% change").font = F(10, True)
            ws.cell(avg + 2, 3, f"=IFERROR(ROUND((C{avg}-C{prev_avg})/C{prev_avg}*100,2),\"\")")
        ws.cell(hdr + 7, 6, "Daily fluid intake (L)").font = F(); ws.cell(avg, 6, "Daily caffeine intake").font = F()
        ws.cell(avg + 1, 6, "Additional notes").font = F()
        for cc in (6, 7, 8):
            ws.cell(hdr + 7, cc).border = B_THIN; ws.cell(avg, cc).border = B_THIN; ws.cell(avg + 1, cc).border = B_THIN
        mr = avg + 4
        for i, (lab, col) in enumerate((("Calories", "E"), ("Protein", "F"), ("Carbs", "G"), ("Fat", "H"))):
            ws.cell(mr + i, 2, lab).font = F()
            for d_i, day in enumerate(spec["days"][:2]):
                if i == 0:
                    ws.cell(mr - 1, 3 + d_i, day["sheet"][:12]).font = F(9, True)
                ws.cell(mr + i, 3 + d_i, f"='{day['sheet'][:31]}'!{col}4").number_format = "0"
        ws.cell(mr, 6, "Changes made").font = F(10, True)
        ws.merge_cells(start_row=mr + 1, start_column=6, end_row=mr + 4, end_column=8)
        ws.cell(mr + 1, 6).alignment = Alignment(wrap_text=True, vertical="top")
        prev_avg = avg
        r = mr + 7

    # ---------------- Food list ----------------
    ws = wb.create_sheet("Food list")
    for col, wdt in zip("ABCDEFGH", (2.7, 8, 6, 44, 9, 9, 9, 9)):
        ws.column_dimensions[col].width = wdt
    for j, h in enumerate(("Qty", "", "Food", "Calories", "Protein", "Carbs", "Fat")):
        cell = ws.cell(2, 2 + j, h); cell.font = F(10, True); cell.fill = BLUE
    r = 3; lastcat = None
    for key, (n, u, basis, k, p, cc_, f, cat) in FOODS.items():
        if key == "free_veg":
            continue
        if cat != lastcat:
            r += 1; ws.cell(r, 2, cat).font = F(11, True); r += 1; lastcat = cat
        ws.cell(r, 2, basis); ws.cell(r, 3, u if u in ("g", "ml") else u); ws.cell(r, 4, n)
        for j, v in enumerate((k, p, cc_, f)):
            cell = ws.cell(r, 5 + j, f"={v/basis:g}*B{r}" if v else 0); cell.number_format = "0.0"; cell.alignment = C
        r += 1
    ws.cell(r + 1, 2, "Change the Qty column to check any amount. Meat, rice, pasta, oats and potato are weighed raw/dry.").font = F(10)

    # ---------------- Free veg + 100 cal fruit ----------------
    ws = wb.create_sheet("Free veg + 100cal fruit")
    for col, wdt in zip("ABCD", (2.7, 26, 26, 32)):
        ws.column_dimensions[col].width = wdt
    ws["B2"] = "FREE VEGGIES, HERBS & SPICES + 100 CALORIE FRUITS"; ws["B2"].font = F(14, True)
    veg = ["Asparagus", "Broccoli", "Broccolini", "Cauliflower", "Cabbage", "Capsicum", "Celery", "Cucumber", "Zucchini", "Eggplant",
           "Lettuce", "Spinach", "Asian greens (bok choy, choy sum)", "Mushroom", "Onion", "Leek", "Fennel", "Tomato", "Bean sprouts",
           "Snow peas", "Sugar snap peas", "Green beans", "Radish", "Squash", "Carrot", "Silverbeet", "Kale", "Rocket"]
    herbs = ["Coriander", "Chilli", "Chives", "Cinnamon", "Basil", "Bay leaf", "Garlic", "Ginger", "Lemongrass", "Mint", "Oregano",
             "Paprika", "Parsley", "Rosemary", "Thyme", "Turmeric", "Cumin", "Pepper", "Salt", "Lemon/lime juice", "Mustard", "Soy sauce (small splash)"]
    fruit = ["1 medium banana (~100g flesh)", "1 large apple (~200g)", "1 large pear (~200g)", "2 small oranges (~250g)", "2 kiwifruit (~160g)",
             "2 small peaches (~250g)", "2 medium plums (~250g)", "2-3 mandarins (~250g)", "Mango cheeks (~170g)", "Pineapple (~250g)",
             "Watermelon (~300g)", "Rockmelon (~400g)", "Strawberries (~400g)", "Blueberries (~200g)", "Raspberries (~200g)",
             "Frozen mixed berries (~200g)", "Grapes (~150g)"]
    for j, (h, lst) in enumerate((("Free vegetables", veg), ("Free herbs & spices", herbs), ("100 calorie fruits", fruit))):
        ws.cell(4, 2 + j, h).font = F(11, True); ws.cell(4, 2 + j).fill = BLUE
        for i, v in enumerate(lst):
            ws.cell(5 + i, 2 + j, v).font = F(10)
    ws.cell(5 + len(veg) + 1, 2, "Free veggies are not counted in the macros (about 30-40 cal per 150g). Potato, sweet potato, corn, peas and pumpkin are NOT free.").font = F(10)

    from openpyxl.worksheet.properties import PageSetupProperties
    for sh in wb.worksheets:
        sh.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        sh.page_setup.fitToWidth = 1; sh.page_setup.fitToHeight = 0
        sh.page_setup.orientation = "portrait"; sh.page_setup.paperSize = sh.PAPERSIZE_A4
        sh.sheet_view.showGridLines = False
    for rr in range(6, 17):
        wb["Cover"].cell(rr, 6).alignment = Alignment(horizontal="left")
    wb.save(out)

def check(spec):
    t = spec.get("targets", {})
    print(f"{'Day':36} {'kcal':>6} {'P':>5} {'C':>5} {'F':>5}")
    ok = True
    for day in spec["days"]:
        tk = tp = tc = tf = 0
        for meal in day["meals"]:
            mk = mp = mc = mf = 0
            for it in meal["items"]:
                k, p, c, f = macros(it["food"], it["qty"]); mk += k; mp += p; mc += c; mf += f
            print(f"  {meal['name'][:34]:34} {mk:6.0f} {mp:5.0f} {mc:5.0f} {mf:5.0f}")
            tk += mk; tp += mp; tc += mc; tf += mf
        dt = day.get("targets", t)
        print(f"{day['sheet'][:36]:36} {tk:6.0f} {tp:5.0f} {tc:5.0f} {tf:5.0f}")
        if dt:
            print(f"{'  target':36} {dt['calories']:6} {dt['protein']:5} {dt['carbs']:5} {dt['fat']:5}")
        if dt:
            dk = (tk - dt["calories"]) / dt["calories"] * 100
            flags = []
            if abs(tk - dt["calories"]) > min(0.05 * dt["calories"], 100): flags.append(f"calories off {tk-dt['calories']:+.0f} kcal (max +/-{min(0.05*dt['calories'],100):.0f})")
            if abs(tp - dt["protein"]) > 10: flags.append(f"protein off {tp-dt['protein']:+.0f}g")
            if abs(tc - dt["carbs"]) > max(15, 0.1 * dt["carbs"]): flags.append(f"carbs off {tc-dt['carbs']:+.0f}g")
            if abs(tf - dt["fat"]) > max(8, 0.1 * dt["fat"]): flags.append(f"fat off {tf-dt['fat']:+.0f}g")
            print("   -> " + ("OK within tolerance" if not flags else "ADJUST: " + "; ".join(flags)))
            ok = ok and not flags
    return ok

# ---------------------------------------------------------------------------
# CHECK-INS (r14 to r23): the lifter's own weekly check-ins, read against the adjustment rules.
# The input is JSON (see SKILL.md section 7); the numbers come from here, never from prose.
# ---------------------------------------------------------------------------
GOAL_KIND = {"fat_loss": "cut", "lean_gain": "gain", "maintain": "maintain", "recomp": "maintain",
             "performance": "maintain"}


def _mean(xs, nd=2):
    vals = [float(x) for x in (xs or []) if x is not None and x != ""]
    return (round(sum(vals) / len(vals), nd), len(vals)) if vals else (None, 0)


def _pct(new, old):
    return round((new - old) / old * 100, 2) if new is not None and old else None


def checkin(data):
    """Weekly means, the trend and the options the r14 to r23 rules list. The lifter decides."""
    stops = hard_stops(data.get("age"), data.get("flags", []))
    if stops:
        raise SystemExit("check-in refused: hard stop (" + ", ".join(stops) + "). " + REFERRAL)
    goal = data.get("goal", "")
    if goal not in GOAL_KIND:
        raise SystemExit("goal must be fat_loss | lean_gain | maintain | recomp | performance")
    kind = GOAL_KIND[goal]
    if data.get("in_season"):
        kind = "maintain"                                          # r20, r24: in season sits at maintenance
    as_of = datetime.date.fromisoformat(data["as_of"])
    weeks = []
    for i, w in enumerate(data.get("weeks", [])):
        mean, n = _mean(w.get("weights"))
        steps, sn = _mean(w.get("steps"), 0)
        weeks.append({"week": i + 1, "start": w.get("start"), "mean": mean, "weigh_ins": n,
                      "steps": steps, "steps_days": sn, "pct_change": None,
                      "no_decision": w.get("no_decision") or None})
    for a, b in zip(weeks, weeks[1:]):
        b["pct_change"] = _pct(b["mean"], a["mean"])               # aligned 7-day means (r14)
    if not weeks:
        raise SystemExit("no weeks given")
    cur, raw = weeks[-1], data["weeks"][-1]
    two = weeks[-3] if len(weeks) >= 3 else None
    trend = (round(_pct(cur["mean"], two["mean"]) / 2, 2)          # r14: 14-day slope, %/week
             if two and cur["mean"] is not None and two["mean"] is not None else None)
    rates = [w["pct_change"] for w in weeks[-2:] if w["pct_change"] is not None] if len(weeks) >= 3 else []
    start = datetime.date.fromisoformat(data["plan_start"])
    since = max(start, datetime.date.fromisoformat(data["last_change"])) if data.get("last_change") else start
    prior_steps = [w["steps"] for w in weeks[:-1] if w["steps"] is not None][-4:]
    baseline = round(sum(prior_steps) / len(prior_steps)) if prior_steps else None
    steps_vs = _pct(cur["steps"], baseline) if baseline else None
    no_decision = [w["no_decision"] for w in weeks[-2:] if w["no_decision"]]
    if len(weeks) == 1:
        no_decision.append("the first week of the plan")
    adherence = raw.get("adherence")
    gates = {
        "wait": "pass" if (as_of - since).days >= 14 else "fail",                                    # r14
        "adherence": None if adherence is None else ("pass" if adherence >= 4 else "fail"),           # r15
        "weigh_ins": "pass" if cur["weigh_ins"] >= 5 else "fail",                                     # r15
        "off_plan_meals": "fail" if (raw.get("off_plan_meals") or 0) >= 2 else "pass",                # r15
        "steps": None if steps_vs is None else ("fail" if steps_vs < -15 else "pass"),                 # r16
        "no_decision": ", ".join(no_decision) or None,                                                # r21
    }
    gates_ok = (gates["wait"] == "pass" and gates["adherence"] == "pass" and gates["weigh_ins"] == "pass"
                and gates["off_plan_meals"] == "pass" and not gates["no_decision"])
    options, rules = [], []

    def opt(rule, change):
        rules.append(rule)
        options.append({"rule": rule, "change": change})

    red = [str(x) for x in data.get("red_flags", []) if str(x).strip()]
    if red:                                                        # r23
        opt("r23", "move to maintenance calories and see a doctor (" + ", ".join(red) + "); no deficit until cleared")
    elif "fail" in (gates["adherence"], gates["weigh_ins"], gates["off_plan_meals"]):
        opt("r15", "no calorie change; fix the process first (meal prep, swaps, weekend structure, alcohol)")
    if not red and len(rates) == 2:
        lean = bool(data.get("lean_or_performance"))
        if kind == "cut":
            if all(r > -0.5 for r in rates) and gates_ok:
                if gates["steps"] == "fail":
                    opt("r16", "steps fell more than 15% from baseline: restore steps before any food cut")
                opt("r17", "5-10% fewer calories from carbohydrate or fat (protein unchanged; fat at least 20% of "
                           "energy and 0.5 g/kg; carbohydrate at least 3 g/kg for field-sport and ADF), or about "
                           "2,000 more steps a day")
            fast = -0.7 if lean else -1.0
            low = [k for k in ("sleep", "energy") if raw.get(k) is not None and raw.get(k) <= 2]
            if all(r < fast for r in rates) or (raw.get("performance_falling") and low):
                opt("r18", "5-10% more calories, carbohydrate first")
            weeks_in = (as_of - start).days // 7
            if weeks_in >= 8:
                opt("r22", f"a cut of {weeks_in} weeks: a 1-2 week diet break at maintenance is due every 3-6 weeks "
                           "in a cut longer than 8-12 weeks; expect a 0.5-1 kg scale rise, left out of the rate")
            elif all(raw.get(k) is not None and raw.get(k) <= 2 for k in ("hunger", "adherence", "training")):
                opt("r22", "hunger, adherence and training all 2/5 or less: an early 1-2 week diet break at "
                           "maintenance; expect a 0.5-1 kg scale rise, left out of the rate")
        elif kind == "gain":
            hi = 0.75 if data.get("status") == "novice" else 0.5
            if all(r < 0.25 for r in rates) and gates_ok:
                opt("r19", "+100-200 kcal, mostly carbohydrate")
            elif all(r > hi for r in rates) and gates_ok:
                opt("r19", "-100-200 kcal")
    if not red and kind == "maintain" and len(weeks) >= 3:
        ref = weeks[-4] if len(weeks) >= 4 else weeks[-3]
        drift = _pct(cur["mean"], ref["mean"]) if ref["mean"] is not None else None
        if drift is not None and abs(drift) > 1.0 and gates_ok:
            opt("r20", ("-" if drift > 0 else "+") + f"100-200 kcal (weight drifted {drift:+.2f}% over "
                       f"{len(weeks) - ref['week']} weeks)")
    if not options:
        if len(rates) < 2:
            options.append({"rule": "r14", "change": "no calorie change yet: a decision needs 3 weeks of weigh-ins "
                                                     "(a 14-day slope)"})
        elif not gates_ok:
            options.append({"rule": "r14", "change": "no calorie change this week: see the gates"})
        else:
            options.append({"rule": None, "change": "no change: the trend is within the rules"})
    return {"goal": goal, "weeks": weeks, "trend_pct_week": trend, "weekly_rates": rates,
            "steps_baseline": baseline, "steps_vs_baseline_pct": steps_vs, "gates": gates,
            "rules": sorted(set(rules)), "options": options,
            "recheck": (as_of + datetime.timedelta(days=14)).isoformat(),
            "note": "One change per 2 weeks (r14). The lifter decides which option, if any, to take."}


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    a = sub.add_parser("targets")
    a.add_argument("--sex", required=True); a.add_argument("--age", type=float, required=True)
    hg = a.add_mutually_exclusive_group(required=True)
    hg.add_argument("--height", type=float, help="cm"); hg.add_argument("--height-in", type=float, help="inches")
    wg = a.add_mutually_exclusive_group(required=True)
    wg.add_argument("--weight", type=float, help="kg"); wg.add_argument("--weight-lb", type=float, help="pounds")
    a.add_argument("--bf", type=float); a.add_argument("--activity", required=True, choices=list(ACTIVITY))
    a.add_argument("--goal", required=True); a.add_argument("--formula", default="mifflin", choices=["mifflin", "cunningham", "harris", "katch"])
    a.add_argument("--sport", default="none", choices=["none", "rugby", "field", "adf", "powerlifting", "other"])
    a.add_argument("--season", choices=["off", "pre", "in"]); a.add_argument("--status", default="intermediate", choices=["novice", "intermediate", "advanced"])
    a.add_argument("--exercise-kcal", type=float, help="average daily training energy, for energy availability")
    a.add_argument("--flags", default="", help="declared hard stops, comma separated (r1)")
    a.add_argument("--peak", action="store_true", help="allow fat at 15-20%% of energy (brief, men, peak phase)")
    a.add_argument("--rate", type=float); a.add_argument("--surplus", type=float)
    a.add_argument("--protein", type=float, help="g/kg override"); a.add_argument("--fatpct", type=float, help="0.20-0.35")
    a.add_argument("--rest-days", type=int, default=0, help="rest days per week, prints training/rest split if >= 2")
    s = sub.add_parser("swap"); s.add_argument("food"); s.add_argument("qty", type=float); s.add_argument("to", nargs="+")
    s.add_argument("--basis", default="auto", choices=["auto", "protein", "carbs", "kcal"])
    fo = sub.add_parser("foods"); fo.add_argument("category", nargs="?")
    me = sub.add_parser("meal"); me.add_argument("items", nargs="+", help="food_key:qty")
    b = sub.add_parser("build"); b.add_argument("spec"); b.add_argument("out")
    ck = sub.add_parser("check"); ck.add_argument("spec")
    ci = sub.add_parser("checkin"); ci.add_argument("checkins")
    args = ap.parse_args()
    split = lambda t: [x.strip() for x in t.split(",") if x.strip()]
    if args.cmd == "targets":
        h = args.height if args.height is not None else round(args.height_in * 2.54, 1)
        w = args.weight if args.weight is not None else round(args.weight_lb * 0.45359237, 1)
        print(json.dumps(targets(args.sex, args.age, h, w, args.activity, args.goal, args.bf, args.formula,
                                 args.rate, args.surplus, args.protein, args.fatpct, args.rest_days, sport=args.sport,
                                 season=args.season, status=args.status, exercise_kcal=args.exercise_kcal,
                                 flags=split(args.flags), peak=args.peak), indent=1))
    elif args.cmd == "meal":
        T = [0, 0, 0, 0]
        for it in args.items:
            key, q = it.rsplit(":", 1); m = macros(key, float(q)); T = [x + y for x, y in zip(T, m)]
            print(f"{fmt_qty(key, float(q)):55} {m[0]:5.0f} kcal  P{m[1]:4.0f}  C{m[2]:4.0f}  F{m[3]:4.0f}")
        print(f"{'TOTAL':55} {T[0]:5.0f} kcal  P{T[1]:4.0f}  C{T[2]:4.0f}  F{T[3]:4.0f}")
    elif args.cmd == "swap":
        for to in args.to:
            q = swap_qty(args.food, args.qty, to, args.basis)
            k, p, c, f = macros(to, q)
            ok_ = macros(args.food, args.qty)[0]
            note = ""
            if ok_ and abs(k - ok_) / ok_ > 0.15: note += f"  (kcal {100*(k-ok_)/ok_:+.0f}%)"
            if food(to)[1] in ("g", "ml") and q > 600: note += "  (big portion: split it, e.g. half rice half potato)"
            print(f"{fmt_qty(to, q):55} {k:5.0f} kcal  P{p:4.0f}  C{c:4.0f}  F{f:4.0f}{note}")
        k, p, c, f = macros(args.food, args.qty)
        print(f"(original) {fmt_qty(args.food, args.qty):44} {k:5.0f} kcal  P{p:4.0f}  C{c:4.0f}  F{f:4.0f}")
    elif args.cmd == "foods":
        for key, v in FOODS.items():
            if not args.category or v[7].lower().startswith(args.category.lower()):
                print(f"{key:22} {v[0]:45} per {v[2]} {v[1]:5} {v[3]:>5} kcal P{v[4]} C{v[5]} F{v[6]}")
    elif args.cmd == "check":
        spec = json.load(open(args.spec)); check(spec)
    elif args.cmd == "checkin":
        print(json.dumps(checkin(json.load(open(args.checkins))), indent=1))
    elif args.cmd == "build":
        spec = json.load(open(args.spec))
        gate(spec)
        ok = check(spec)
        build(spec, args.out)
        print(f"\nSaved {args.out}" + ("" if ok else "  (WARNING: some days outside tolerance)"))

if __name__ == "__main__":
    main()
